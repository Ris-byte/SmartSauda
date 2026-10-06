from __future__ import annotations

import json
import math
import re
from pathlib import Path
from xml.sax.saxutils import escape

from PIL import Image as PILImage, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen.canvas import Canvas
from reportlab.platypus import Image, Paragraph, Table, TableStyle


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'docs/SmartSauda_Final_Report.md'
OUT = ROOT / 'docs'
ASSETS = OUT / 'report-assets'
FONT_DIR = Path('C:/Windows/Fonts')
PAGE_WIDTH, PAGE_HEIGHT = 595.32, 841.92
LEFT, RIGHT, WIDTH = 108.02, 72, 415.30
AUTHORS = 'Prabin Thanet; Rishab Magar; Rajan Shah'


def diagram(name, nodes, edges, size=(1200, 700)):
    canvas = PILImage.new('RGB', size, 'white')
    draw = ImageDraw.Draw(canvas)
    font = ImageFont.truetype(str(FONT_DIR / 'times.ttf'), 28)
    bold = ImageFont.truetype(str(FONT_DIR / 'timesbd.ttf'), 29)
    label_font = ImageFont.truetype(str(FONT_DIR / 'times.ttf'), 24)

    def centre(node):
        left, top, right, bottom = nodes[node][0]
        return ((left + right) / 2, (top + bottom) / 2)

    def boundary(node, toward):
        bounds, _, kind = nodes[node]
        centre_x, centre_y = centre(node)
        delta_x, delta_y = toward[0] - centre_x, toward[1] - centre_y
        half_width = (bounds[2] - bounds[0]) / 2
        half_height = (bounds[3] - bounds[1]) / 2
        if kind == 'oval':
            scale = 1 / math.sqrt((delta_x / half_width) ** 2 + (delta_y / half_height) ** 2)
        else:
            scale = min(half_width / abs(delta_x) if delta_x else 1e9,
                        half_height / abs(delta_y) if delta_y else 1e9)
        return centre_x + scale * delta_x, centre_y + scale * delta_y

    for source, target, label in edges:
        bidirectional = label.startswith('<>')
        label = label.removeprefix('<>')
        start = boundary(source, centre(target))
        end = boundary(target, centre(source))
        draw.line([start, end], fill='black', width=3)
        angle = math.atan2(end[1] - start[1], end[0] - start[0])
        arrow = [end, (end[0] - 14 * math.cos(angle - .45), end[1] - 14 * math.sin(angle - .45)),
                 (end[0] - 14 * math.cos(angle + .45), end[1] - 14 * math.sin(angle + .45))]
        draw.polygon(arrow, fill='black')
        if bidirectional:
            reverse_angle = angle + math.pi
            draw.polygon([start, (start[0] - 14 * math.cos(reverse_angle - .45), start[1] - 14 * math.sin(reverse_angle - .45)),
                          (start[0] - 14 * math.cos(reverse_angle + .45), start[1] - 14 * math.sin(reverse_angle + .45))], fill='black')
        if label:
            midpoint = ((start[0] + end[0]) / 2, (start[1] + end[1]) / 2 - 17)
            bounds = draw.textbbox(midpoint, label, font=label_font, anchor='mm')
            draw.rectangle((bounds[0] - 4, bounds[1] - 2, bounds[2] + 4, bounds[3] + 2), fill='white')
            draw.text(midpoint, label, font=label_font, fill='black', anchor='mm')

    for bounds, text, kind in nodes.values():
        if kind == 'oval':
            draw.ellipse(bounds, fill='white', outline='black', width=3)
        elif kind == 'store':
            draw.rectangle(bounds, fill='white', outline='black', width=3)
            draw.line((bounds[0] + 15, bounds[1], bounds[0] + 15, bounds[3]), fill='black', width=2)
        else:
            draw.rectangle(bounds, fill='white', outline='black', width=3)
        if kind == 'entity':
            draw.rectangle((bounds[0] + 2, bounds[1] + 2, bounds[2] - 2, bounds[1] + 43), fill='#dce9f1')
            lines = text.split('\n')
            draw.text(((bounds[0] + bounds[2]) / 2, bounds[1] + 22), lines[0], font=bold, anchor='mm', fill='black')
            draw.multiline_text((bounds[0] + 15, bounds[1] + 55), '\n'.join(lines[1:]), font=font, fill='black', spacing=8)
        else:
            centre_x, centre_y = (bounds[0] + bounds[2]) / 2, (bounds[1] + bounds[3]) / 2
            draw.multiline_text((centre_x, centre_y), text, font=font, fill='black', anchor='mm', align='center', spacing=8)
    path = ASSETS / (name + '.png')
    canvas.save(path, dpi=(220, 220))


def make_diagrams():
    ASSETS.mkdir(parents=True, exist_ok=True)
    diagram('workflow', {
        'requirements': ((350, 20, 850, 100), 'Requirements and data review', 'box'),
        'backend': ((20, 200, 350, 290), 'Backend and database', 'box'),
        'frontend': ((435, 200, 765, 290), 'Frontend design', 'box'),
        'model': ((850, 200, 1180, 290), 'Model development', 'box'),
        'integration': ((350, 400, 850, 490), 'Application integration', 'box'),
        'review': ((350, 580, 850, 670), 'Testing and documentation', 'box'),
    }, [('requirements', 'backend', ''), ('requirements', 'frontend', ''), ('requirements', 'model', ''),
        ('backend', 'integration', ''), ('frontend', 'integration', ''), ('model', 'integration', ''), ('integration', 'review', '')])
    diagram('usecase', {
        'visitor': ((10, 40, 200, 120), 'Visitor', 'box'),
        'user': ((10, 320, 200, 420), 'Registered\nuser', 'box'),
        'admin': ((10, 580, 200, 660), 'Administrator', 'box'),
        'explore': ((385, 20, 825, 120), 'Explore and view catalog', 'oval'),
        'auth': ((700, 170, 1175, 270), 'Register / sign in', 'oval'),
        'predict': ((375, 305, 825, 410), 'Create vehicle estimate', 'oval'),
        'history': ((700, 455, 1175, 555), 'View history / download report', 'oval'),
        'manage': ((385, 585, 825, 685), 'Manage account status', 'oval'),
    }, [('visitor', 'explore', ''), ('visitor', 'auth', ''), ('user', 'auth', ''), ('user', 'predict', ''),
        ('user', 'history', ''), ('admin', 'manage', '')])
    diagram('dfd', {
        'browser': ((30, 40, 330, 130), 'User / browser', 'box'),
        'validate': ((440, 20, 780, 160), '1.0\nValidate request', 'oval'),
        'catalog': ((900, 40, 1180, 130), 'D1 Catalog', 'store'),
        'infer': ((440, 280, 780, 420), '2.0\nEstimate price', 'oval'),
        'models': ((900, 300, 1180, 390), 'Model artifacts', 'store'),
        'result': ((30, 570, 330, 680), 'Result / PDF\nto user', 'box'),
        'retrieve': ((440, 555, 780, 695), '3.0\nSave / retrieve', 'oval'),
        'db': ((900, 575, 1180, 670), 'D2 Predictions', 'store'),
    }, [('browser', 'validate', 'details'), ('catalog', 'validate', 'limits'), ('validate', 'infer', 'valid inputs'),
        ('models', 'infer', 'estimator'), ('infer', 'retrieve', 'result'), ('retrieve', 'db', '<>save / read'),
        ('retrieve', 'result', 'record')])
    diagram('erd', {
        'users': ((395, 25, 810, 265), 'users\nid (PK)\nemail (unique)\npassword_hash\nrole, active', 'entity'),
        'sessions': ((20, 400, 430, 645), 'sessions\ntoken_digest (PK)\nuser_id (FK)\nexpires_at', 'entity'),
        'predictions': ((750, 400, 1180, 680), 'predictions\nid (PK), user_id (FK)\nmodel_version, price\nspecifications, result\nimage, created_at', 'entity'),
    }, [('users', 'sessions', '1 : many'), ('users', 'predictions', '1 : many')])
    diagram('architecture', {
        'ui': ((330, 20, 870, 125), 'React + TypeScript\nBrowser interface', 'box'),
        'api': ((300, 265, 900, 390), 'FastAPI application\nAuth / validation / reports', 'box'),
        'db': ((20, 575, 370, 685), 'PostgreSQL\nAccounts and predictions', 'store'),
        'ml': ((425, 575, 775, 685), 'Category models\nPreprocessing + regression', 'box'),
        'images': ((830, 575, 1180, 685), 'Vehicle images\nCache and fallback', 'box'),
    }, [('ui', 'api', 'HTTP / JSON'), ('api', 'db', 'read / write'), ('api', 'ml', 'inference'), ('api', 'images', 'lookup')])
    diagram('training', {
        'data': ((30, 30, 530, 135), 'Retained vehicle records\n3,316 rows', 'box'),
        'split': ((680, 30, 1180, 135), 'Group-aware partitions\nTraining / validation / test', 'box'),
        'fit': ((680, 290, 1180, 395), 'Fit and tune candidates\nTraining partition only', 'box'),
        'select': ((30, 290, 530, 395), 'Select using validation MAE\nRefit on training + validation', 'box'),
        'test': ((30, 560, 530, 665), 'Evaluate held-out test data\nDo not fit on test records', 'box'),
        'export': ((680, 560, 1180, 665), 'Export fitted model\nMetadata and evaluation', 'box'),
    }, [('data', 'split', ''), ('split', 'fit', ''), ('fit', 'select', ''), ('select', 'test', ''), ('test', 'export', '')])


def parse_page(text):
    lines = text.strip().splitlines()
    blocks = []
    cursor = 0
    while cursor < len(lines):
        line = lines[cursor].strip()
        if not line:
            cursor += 1
            continue
        if line.startswith('#'):
            depth = len(line) - len(line.lstrip('#'))
            blocks.append(('heading', depth, line[depth:].strip()))
        elif line.startswith('{{'):
            blocks.append(('index', line[2:-2]))
        elif line.startswith('!['):
            match = re.match(r'!\[(.*?)\]\((.*?)\)', line)
            blocks.append(('figure', match[1], SOURCE.parent / match[2]))
        elif line.startswith('|'):
            rows = []
            while cursor < len(lines) and lines[cursor].strip().startswith('|'):
                row = [cell.strip() for cell in lines[cursor].strip().strip('|').split('|')]
                if not all(re.fullmatch(r'[:\- ]+', cell) for cell in row):
                    rows.append(row)
                cursor += 1
            blocks.append(('table', rows))
            continue
        elif re.match(r'^Table(?: \d+)?:', line):
            blocks.append(('caption', line))
        elif line.startswith('- '):
            blocks.append(('bullet', line[2:]))
        else:
            paragraph = [line]
            while cursor + 1 < len(lines) and lines[cursor + 1].strip() and not lines[cursor + 1].startswith(('#', '|', '![', '{{', '- ')):
                cursor += 1
                paragraph.append(lines[cursor].strip())
            blocks.append(('text', '\n'.join(paragraph)))
        cursor += 1
    return blocks


def roman(number):
    output = ''
    for value, letters in [(50, 'l'), (40, 'xl'), (10, 'x'), (9, 'ix'), (5, 'v'), (4, 'iv'), (1, 'i')]:
        while number >= value:
            output += letters
            number -= value
    return output


def collect_indexes(pages, body_start):
    contents, figures, tables = [], [], []
    for index, blocks in enumerate(pages):
        page_number = str(index - body_start + 1) if index >= body_start else roman(index)
        for block in blocks:
            if block[0] == 'heading' and index >= body_start and '(Continued)' not in block[2]:
                contents.append((block[2], page_number, block[1]))
            if block[0] == 'figure':
                figures.append((block[1], page_number, 1))
            if block[0] == 'caption' and re.match(r'Table \d+:', block[1]):
                tables.append((block[1], page_number, 1))
    split = math.ceil(len(contents) / 2)
    return {'CONTENTS_1': contents[:split], 'CONTENTS_2': contents[split:], 'FIGURES': figures, 'TABLES': tables}


def register_fonts():
    for name, filename in [('TNR', 'times.ttf'), ('TNR-Bold', 'timesbd.ttf'), ('TNR-Italic', 'timesi.ttf'), ('TNR-BoldItalic', 'timesbi.ttf')]:
        pdfmetrics.registerFont(TTFont(name, str(FONT_DIR / filename)))
    pdfmetrics.registerFontFamily('TNR', normal='TNR', bold='TNR-Bold', italic='TNR-Italic', boldItalic='TNR-BoldItalic')


def markup(text):
    return escape(text.replace('`', '')).replace('\n', '<br/>')


def styles():
    return {
        'text': ParagraphStyle('body', fontName='TNR', fontSize=12, leading=15.84, alignment=TA_JUSTIFY),
        'heading1': ParagraphStyle('chapter', fontName='TNR-Bold', fontSize=16, leading=20, alignment=TA_CENTER),
        'heading2': ParagraphStyle('section', fontName='TNR-Bold', fontSize=14, leading=18, alignment=TA_LEFT),
        'heading3': ParagraphStyle('subsection', fontName='TNR-Bold', fontSize=12, leading=15.84, alignment=TA_LEFT),
        'caption': ParagraphStyle('caption', fontName='TNR-Bold', fontSize=11, leading=14, alignment=TA_CENTER),
        'cell': ParagraphStyle('cell', fontName='TNR', fontSize=11, leading=14, alignment=TA_LEFT, splitLongWords=True),
        'index': ParagraphStyle('index', fontName='TNR', fontSize=11, leading=14, alignment=TA_LEFT),
        'cover': ParagraphStyle('cover', fontName='TNR-Bold', fontSize=14, leading=18, alignment=TA_CENTER),
    }


def table_widths(rows):
    count = len(rows[0])
    header = rows[0][0]
    if header == 'Phase / Task' and count == 3:
        ratios = [.22, .55, .23]
    elif header == 'ID' and count == 4:
        ratios = [.11, .35, .33, .21]
    elif count == 2:
        ratios = [.25, .75]
    elif count == 3:
        ratios = [.23, .31, .46]
    elif count == 4:
        ratios = [.29, .39, .16, .16]
    elif count == 5 and rows[0][1] == 'Selected Model':
        ratios = [.13, .31, .12, .23, .21]
    else:
        ratios = [1 / count] * count
    return [WIDTH * ratio for ratio in ratios]


def build_pdf(pages, body_start, indexes):
    register_fonts()
    formatting = styles()
    path = OUT / 'SmartSauda_Final_Report.pdf'
    canvas = Canvas(str(path), pagesize=(PAGE_WIDTH, PAGE_HEIGHT))
    canvas.setTitle('SmartSauda: Vehicle Valuation and Resale Estimation System')
    canvas.setAuthor(AUTHORS)
    canvas.setSubject('Project Report')
    canvas.setCreator('')
    canvas.setProducer('')
    measurements = []
    for index, blocks in enumerate(pages):
        top = 72 if index < body_start else 90
        position = PAGE_HEIGHT - top
        cover = index == 0
        for block in blocks:
            kind = block[0]
            if kind == 'index':
                for title, number, depth in indexes[block[1]]:
                    indent = 10 * (depth - 1)
                    font_size = 10.5 if block[1].startswith('CONTENTS') else 11
                    available = WIDTH - indent - 25
                    while pdfmetrics.stringWidth(title, 'TNR', font_size) > available and font_size > 8.5:
                        font_size -= .25
                    canvas.setFont('TNR-Bold' if depth == 1 else 'TNR', font_size)
                    position -= 14.5
                    canvas.drawString(LEFT + indent, position, title)
                    canvas.setFont('TNR', 11)
                    canvas.drawRightString(LEFT + WIDTH, position, number)
                    title_width = pdfmetrics.stringWidth(title, 'TNR-Bold' if depth == 1 else 'TNR', font_size)
                    dots_start = LEFT + indent + title_width + 5
                    dots_end = LEFT + WIDTH - 24
                    if dots_end > dots_start:
                        canvas.setDash(1, 2)
                        canvas.setLineWidth(.3)
                        canvas.line(dots_start, position + 2, dots_end, position + 2)
                        canvas.setDash()
                continue
            if kind == 'figure':
                with PILImage.open(block[2]) as source_image:
                    aspect = source_image.height / source_image.width
                image_height = min(WIDTH * aspect, 140 if block[2].stem == 'workflow' else 270)
                image_width = image_height / aspect
                figure = Image(str(block[2]), width=image_width, height=image_height)
                figure.drawOn(canvas, LEFT + (WIDTH - image_width) / 2, position - image_height)
                position -= image_height + 8
                caption = Paragraph(markup(block[1]), formatting['caption'])
                _, height = caption.wrap(WIDTH, 100)
                caption.drawOn(canvas, LEFT, position - height)
                position -= height + 17
                continue
            if kind == 'table':
                rows = [[Paragraph(('<b>' if row_index == 0 else '') + markup(cell) + ('</b>' if row_index == 0 else ''), formatting['cell']) for cell in row]
                        for row_index, row in enumerate(block[1])]
                table = Table(rows, colWidths=table_widths(block[1]))
                table.setStyle(TableStyle([('GRID', (0, 0), (-1, -1), .45, colors.HexColor('#777777')),
                                           ('VALIGN', (0, 0), (-1, -1), 'TOP'), ('LEFTPADDING', (0, 0), (-1, -1), 5),
                                           ('RIGHTPADDING', (0, 0), (-1, -1), 5), ('TOPPADDING', (0, 0), (-1, -1), 6),
                                           ('BOTTOMPADDING', (0, 0), (-1, -1), 6)]))
                _, height = table.wrap(WIDTH, PAGE_HEIGHT)
                table.drawOn(canvas, LEFT, position - height)
                position -= height + 10
                continue
            content = block[2] if kind == 'heading' else block[1]
            style_name = 'heading' + str(block[1]) if kind == 'heading' else 'caption' if kind == 'caption' else 'text'
            if cover:
                style_name = 'cover'
            if kind == 'bullet':
                content = '\u2022  ' + content
            paragraph = Paragraph(markup(content), formatting[style_name])
            _, height = paragraph.wrap(WIDTH, PAGE_HEIGHT)
            paragraph.drawOn(canvas, LEFT, position - height)
            gap = 18 if kind == 'heading' else 12 if kind == 'caption' else 15.84
            if kind == 'bullet':
                gap = 7
            if cover:
                gap = 20
            position -= height + gap
        measurements.append({'physical_page': index + 1, 'content_bottom': round(position, 2)})
        if position < 63:
            raise ValueError(f'Page {index + 1} overflows: content bottom {position:.1f}')
        if not cover:
            label = str(index - body_start + 1) if index >= body_start else roman(index)
            canvas.setFont('TNR', 12)
            canvas.drawCentredString(LEFT + WIDTH / 2, 56.544, label)
        canvas.showPage()
    canvas.save()
    return measurements


def set_numbering(section, number_format):
    section.footer.is_linked_to_previous = False
    paragraph = section.footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(0)
    element = OxmlElement('w:fldSimple')
    element.set(qn('w:instr'), 'PAGE')
    paragraph._p.append(element)
    numbering = OxmlElement('w:pgNumType')
    numbering.set(qn('w:start'), '1')
    numbering.set(qn('w:fmt'), number_format)
    section._sectPr.append(numbering)


def section_format(section, body=False):
    section.page_width = Pt(PAGE_WIDTH)
    section.page_height = Pt(PAGE_HEIGHT)
    section.left_margin = Pt(LEFT)
    section.right_margin = Pt(RIGHT)
    section.top_margin = Pt(90 if body else 72)
    section.bottom_margin = Pt(75)
    section.footer_distance = Pt(48)


def paragraph_format(paragraph, after=15.84, line=15.84, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY):
    paragraph.alignment = alignment
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(after)
    fmt.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    fmt.line_spacing = Pt(line)
    fmt.widow_control = False


def build_docx(pages, body_start, indexes):
    document = Document()
    document.core_properties.title = 'SmartSauda: Vehicle Valuation and Resale Estimation System'
    document.core_properties.author = AUTHORS
    document.core_properties.last_modified_by = ''
    document.core_properties.subject = 'Project Report'
    document.core_properties.comments = ''
    for style_name in ['Normal', 'Heading 1', 'Heading 2', 'Heading 3', 'Caption']:
        style = document.styles[style_name]
        style.font.name = 'Times New Roman'
        style.font.size = Pt(12)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(15.84)
        font_element = style.element.get_or_add_rPr().find(qn('w:rFonts'))
        for attribute in list(font_element.attrib):
            if 'theme' in attribute.lower():
                del font_element.attrib[attribute]
        for attribute in ['ascii', 'hAnsi', 'eastAsia', 'cs']:
            font_element.set(qn('w:' + attribute), 'Times New Roman')
    section_format(document.sections[0])
    for index, blocks in enumerate(pages):
        if index in (1, body_start):
            section = document.add_section(WD_SECTION_START.NEW_PAGE)
            section_format(section, index == body_start)
            set_numbering(section, 'decimal' if index == body_start else 'lowerRoman')
        elif index:
            paragraph = document.add_paragraph()
            paragraph_format(paragraph, after=0, line=1)
            paragraph.add_run().add_break(WD_BREAK.PAGE)
        for block in blocks:
            kind = block[0]
            if kind == 'index':
                for title, number, depth in indexes[block[1]]:
                    paragraph = document.add_paragraph()
                    paragraph_format(paragraph, after=0, line=14.5, alignment=WD_ALIGN_PARAGRAPH.LEFT)
                    paragraph.paragraph_format.left_indent = Pt(10 * (depth - 1))
                    paragraph.paragraph_format.tab_stops.add_tab_stop(Pt(WIDTH), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
                    run = paragraph.add_run(title + '\t' + number)
                    run.font.size = Pt(10.5)
                    run.bold = depth == 1
                continue
            if kind == 'figure':
                with PILImage.open(block[2]) as source_image:
                    aspect = source_image.height / source_image.width
                image_height = min(WIDTH * aspect, 140 if block[2].stem == 'workflow' else 270)
                paragraph = document.add_paragraph()
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.space_after = Pt(8)
                paragraph.paragraph_format.keep_with_next = True
                paragraph.add_run().add_picture(str(block[2]), width=Pt(image_height / aspect), height=Pt(image_height))
                paragraph = document.add_paragraph(block[1])
                paragraph_format(paragraph, after=17, line=14, alignment=WD_ALIGN_PARAGRAPH.CENTER)
                for run in paragraph.runs:
                    run.bold = True
                    run.font.size = Pt(11)
                continue
            if kind == 'table':
                widths = table_widths(block[1])
                table = document.add_table(rows=0, cols=len(widths))
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                table.autofit = False
                table.style = 'Table Grid'
                for column, width in zip(table.columns, widths):
                    column.width = Pt(width)
                for row_index, row in enumerate(block[1]):
                    cells = table.add_row().cells
                    for cell, text, width in zip(cells, row, widths):
                        cell.width = Pt(width)
                        cell.text = text
                        paragraph_format(cell.paragraphs[0], after=0, line=14, alignment=WD_ALIGN_PARAGRAPH.LEFT)
                        for run in cell.paragraphs[0].runs:
                            run.font.size = Pt(11)
                            run.bold = row_index == 0
                        margins = OxmlElement('w:tcMar')
                        for tag, value in [('top', 100), ('bottom', 100), ('left', 100), ('right', 100)]:
                            child = OxmlElement('w:' + tag)
                            child.set(qn('w:w'), str(value))
                            child.set(qn('w:type'), 'dxa')
                            margins.append(child)
                        cell._tc.get_or_add_tcPr().append(margins)
                    no_split = OxmlElement('w:cantSplit')
                    table.rows[-1]._tr.get_or_add_trPr().append(no_split)
                paragraph = document.add_paragraph()
                paragraph_format(paragraph, after=0, line=10)
                continue
            text = block[2] if kind == 'heading' else block[1]
            if kind == 'bullet':
                text = '\u2022  ' + text
            paragraph = document.add_paragraph(text.replace('`', ''))
            paragraph_format(paragraph)
            if kind == 'heading':
                depth = block[1]
                paragraph.style = document.styles['Heading ' + str(depth)]
                paragraph_format(paragraph, after=18, line=20 if depth == 1 else 18 if depth == 2 else 15.84,
                                 alignment=WD_ALIGN_PARAGRAPH.CENTER if depth == 1 else WD_ALIGN_PARAGRAPH.LEFT)
                for run in paragraph.runs:
                    run.bold = True
                    run.font.size = Pt(16 if depth == 1 else 14 if depth == 2 else 12)
            elif kind == 'caption':
                paragraph_format(paragraph, after=12, line=14, alignment=WD_ALIGN_PARAGRAPH.CENTER)
                for run in paragraph.runs:
                    run.bold = True
                    run.font.size = Pt(11)
            elif kind == 'bullet':
                paragraph.paragraph_format.space_after = Pt(7)
            if index == 0:
                paragraph_format(paragraph, after=20, line=18, alignment=WD_ALIGN_PARAGRAPH.CENTER)
                for run in paragraph.runs:
                    run.bold = True
                    run.font.size = Pt(14)
    document.save(OUT / 'SmartSauda_Final_Report.docx')


def main():
    make_diagrams()
    pages = [parse_page(page) for page in SOURCE.read_text(encoding='utf-8').split('<!-- page -->')]
    body_start = next(index for index, blocks in enumerate(pages) if blocks[0][0] == 'heading' and blocks[0][2].startswith('CHAPTER 1:'))
    indexes = collect_indexes(pages, body_start)
    measurements = build_pdf(pages, body_start, indexes)
    build_docx(pages, body_start, indexes)
    (ASSETS / 'layout-check.json').write_text(json.dumps({'pages': len(pages), 'body_start': body_start + 1,
                                                      'indexes': indexes, 'measurements': measurements}, indent=2), encoding='utf-8')
    print(f'Created PDF and Word report: {len(pages)} pages, {len(indexes["FIGURES"])} figures.')


if __name__ == '__main__':
    main()
